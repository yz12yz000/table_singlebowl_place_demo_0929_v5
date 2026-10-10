from pathlib import Path
import json, re, zipfile
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[1]
A=OUT/'assets'
prs=Presentation()
prs.slide_width=Inches(13.333333)
prs.slide_height=Inches(7.5)
prs.core_properties.title='Geometry-Constrained Bowl Manipulation'
prs.core_properties.subject='System design, geometric principles, and experimental progress'
prs.core_properties.author='Robotics Project'
prs.core_properties.keywords='RGB-D, radial rim grasping, taught placement, Chinese speaker notes'
BG='F5F7F8'; DARK='112D39'; BLUE='087C99'; TEAL='117D70'; ORANGE='C98125'; GREY='526773'; LIGHT='DCE5E9'; WHITE='FFFFFF'
FONT='Liberation Sans'
records=[]

def rgb(c):return RGBColor.from_string(c)
def rect(s,x,y,w,h,c,line=None):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=rgb(c)
 sh.line.fill.background() if not line else None
 if line:sh.line.color.rgb=rgb(line)
 return sh

def text(s,txt,x,y,w,h,size=18,color=DARK,bold=False,align=None,font=FONT):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
 tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(txt.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
  p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
  p.space_before=Pt(0);p.space_after=Pt(6);p.line_spacing=1.12
  if align is not None:p.alignment=align
 return sh

def img(s,name,x,y,w,h,mode='contain',crop=None):
 path=A/name
 iw,ih=Image.open(path).size
 if crop:
  cl,ct,cr,cb=crop;aspect=iw*(1-cl-cr)/(ih*(1-ct-cb))
  dw=min(w,h*aspect);dh=dw/aspect
  sh=s.shapes.add_picture(str(path),Inches(x+(w-dw)/2),Inches(y+(h-dh)/2),width=Inches(dw),height=Inches(dh))
  sh.crop_left=cl;sh.crop_top=ct;sh.crop_right=cr;sh.crop_bottom=cb
 elif mode=='cover':
  aspect=iw/ih;target=w/h
  sh=s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
  if aspect>target:sh.crop_left=sh.crop_right=(1-target/aspect)/2
  else:sh.crop_top=sh.crop_bottom=(1-aspect/target)/2
 else:
  r=min(w/iw,h/ih);dw=iw*r;dh=ih*r
  sh=s.shapes.add_picture(str(path), Inches(x+(w-dw)/2), Inches(y+(h-dh)/2),width=Inches(dw),height=Inches(dh))
 return sh

def label(s,txt,x,y,w=3,color=BLUE):
 text(s,txt.upper(),x,y,w,.25,10.5,color,True)

def points(s,rows,x,y,w,size=18,gap=.94):
 for i,(head,body) in enumerate(rows):
  yy=y+i*gap
  rect(s,x,yy+.055,.045,.23,BLUE)
  text(s,head,x+.18,yy,w-.18,.35,size,DARK,True)
  text(s,body,x+.18,yy+.39,w-.18,gap-.39,15.5,GREY)

def base(title,kicker,source):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
 rect(s,.48,.40,.08,.27,BLUE)
 label(s,kicker,.70,.40,11)
 text(s,title,.65,.89,12.05,.80,29,DARK,True)
 rect(s,.65,7.04,12.02,.012,LIGHT)
 text(s,source,.65,7.13,10.95,.19,8.3,GREY)
 text(s,f'{len(prs.slides):02d} / 12',11.70,7.10,.95,.23,10,GREY,align=PP_ALIGN.RIGHT)
 return s

def caption(s,txt,x,y,w,color=GREY):text(s,txt,x,y,w,.48,10,color)
def note(s,title,body,sources):
 full='【中文演讲备注】\n'+body+'\n\n【资料来源】\n'+'\n'.join(sources)
 s.notes_slide.notes_text_frame.text=full
 for p in s.notes_slide.notes_text_frame.paragraphs:
  p.font.name='Noto Sans CJK SC';p.font.size=Pt(12)
 records.append({'slide':len(prs.slides),'title':title,'notes_zh':body,'sources':sources})

# 1: original photograph, full-height cover
s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(DARK)
img(s,'lifted_white.jpg',5.08,0,8.26,7.5,'cover')
rect(s,0,0,5.33,7.5,DARK)
label(s,'ROBOTICS / RESEARCH PROGRESS',.60,.65,4.25,'8ACFE1')
text(s,'Geometry-\nConstrained\nBowl Manipulation',.60,1.56,4.50,2.40,33,WHITE,True)
text(s,'From visual rim grasping\nto repeatable placement',.62,4.12,4.25,1.0,21,'D2E5EB')
rect(s,.62,5.47,.68,.055,'55BCD4')
text(s,'RGB-D perception\nRadial contact geometry\nTaught placement',.62,5.77,4.0,1.03,15,'C8DCE3')
text(s,'02 OCT 2026',.62,7.01,3,.22,10,'C8DCE3')
rect(s,8.37,6.68,4.47,.48,DARK)
text(s,'Original hardware image | 23 Sep 2026',8.57,6.81,4.10,.22,10,WHITE)
note(s,'Geometry-Constrained Bowl Manipulation',
'本项目围绕规则圆碗的视觉抓取和指定位置放置展开。除了实现“看到碗、夹起来、放下去”，本阶段重点研究抓取方位变化时，如何尽量保持碗与夹爪之间的相对关系。汇报分为几何原理、完整任务执行和实验证据三个部分。封面使用 9 月 23 日的真实抬碗照片，用于展示硬件平台和抓取场景；这张早期照片不作为后续径向方法的对比验证。',
['docs/圆碗径向抓取修改_20260928.md','debug_output/planning_diagnosis_20260923/white_lifted.jpg'])

# 2: illustrative contact configurations
s=base('A fixed tool pose does not fix the bowl center','DESIGN QUESTION','Principle illustration; no experimental measurements are represented.')
img(s,'contact_geometry.png',.65,1.95,8.00,4.72)
label(s,'ILLUSTRATIVE RENDERING',.85,2.08,4,ORANGE)
points(s,[('Contact geometry matters','Different contacts change the bowl-to-gripper offset.'),('Teaching controls the tool','A fixed release pose does not directly constrain the bowl center.'),('Design objective','Make the grasp relationship more consistent.')],9.00,2.16,3.55,gap=1.40)
caption(s,'AI-generated contact illustration | Same bowl type; different rim engagement',.83,6.58,8)
note(s,'A Fixed Tool Pose Does Not Fix the Bowl Center',
'首先解释为什么只把所有释放点设置成相同 RPY，并不能保证每次落在同一个位置。示教和运动控制约束的是工具位姿，而碗中心还取决于夹爪夹住了碗的哪个位置、以什么方向和深度夹住。接触关系不同，工具到碗中心的偏移就可能不同。这里使用的是 AI 生成的接触示意，不是两次实际实验，也不表示测得了某个偏移量。我们要做的是先稳定抓取关系，再研究放置误差。',
['docs/圆碗径向抓取修改_20260928.md','assets/contact_geometry.png (AI-generated illustration)'])

# 3: real gripper + axis identification
s=base('Radial alignment with fixed finger assignment','CONTRIBUTION 01 / GRASP GEOMETRY','Original user reference images; the physical left finger was identified as tool +Y.')
img(s,'left_finger.jpg',.65,1.87,4.45,4.74,crop=(.38,.38,.22,.38))
img(s,'finger_axis.png',5.17,1.88,3.58,3.62,crop=(.50,.30,.28,.03))
caption(s,'Physical finger identification\nCircled finger = left',.74,6.57,4.2)
caption(s,'RViz frame check\nGreen axis = tool +Y',5.37,5.61,3.2)
points(s,[('Left outside; right inside','Preserve the identity of each physical finger.'),('Radial closing direction','Tool +Y points outward; −Y points toward the bowl center.'),('Consistent engagement','Keep grasp height and insertion depth consistent.')],9.02,1.99,3.65,gap=1.55)
note(s,'Radial Alignment with Fixed Finger Assignment',
'核心约束包括两个部分：夹爪闭合轴沿圆心到碗沿的径向，同一个物理手指始终在碗外。用户已经通过照片和 RViz 确认左指对应工具正 Y，因此当前配置保持左外右内。这样可以减少不同抓取方位以及方向翻转造成的相对关系变化。圆碗的旋转对称性并不意味着带工具偏置和线缆的实际夹爪可以随意翻转。固定抓取高度、深度和避免滑动同样是这一设计成立的前提。',
['docs/left_finger_reference_20260928.jpeg','docs/left_finger_positive_y_rviz_20260928.png','src/rim_locator/rim_locator/rim_detection.py'])

# 4: actual point-cloud evidence
s=base('Global shape guides, local geometry validates','CONTRIBUTION 02 / OBSERVED CONTACT SUPPORT','Original point-cloud replay visualization | circle_radial_replay_final_20260928')
img(s,'radial_geometry.png',.65,1.83,5.15,5.02)
label(s,'OBSERVED GEOMETRY',6.13,1.98,5)
points(s,[('Global shape','Fit the observed rim to estimate center and radial direction.'),('Local contact support','Check wall patches at multiple neighborhood scales.'),('Cross-check the grasp','Compare the wall direction and observed contact with the circle.'),('Preserve measured positions','Do not project grasp contacts onto an idealized rim.')],6.13,2.49,6.35,gap=1.0)
note(s,'Global Shape Guides, Local Geometry Validates',
'这里的设计将整体形状和局部接触分工。圆拟合提供圆心、半径以及有符号的径向方向；局部碗壁搜索验证该位置是否存在可靠的接触证据。实现中会比较不同尺度的局部法线，并检查点数、高度跨度、残差和方向一致性。圆模型通过后，候选还要满足接触半径误差与壁面方向偏差条件。重要的是，输出抓取点保留真实观测位置，不把点移动到拟合圆周上。图来自保存点云的实际算法回放。',
['src/rim_locator/rim_locator/circle_fit.py','src/rim_locator/rim_locator/rim_detection.py','debug_output/circle_radial_replay_final_20260928/radial_grasps.svg'])

# 5: partial arc and fallback evidence
s=base('Use reliable local evidence when the rim is incomplete','METHOD / PARTIAL OBSERVATIONS','Separate live and offline observations from the same diagnostic session | 28 Sep 2026')
img(s,'fallback_detection.png',.65,1.96,5.60,3.50)
img(s,'fallback_geometry.png',7.15,1.80,5.05,4.58)
label(s,'LIVE PLANNING PREVIEW',.77,5.62,6)
text(s,'Approximately 103° coverage; 8 grasp candidates.',.77,5.98,6.36,.39,18,DARK,True)
rect(s,.70,6.52,11.91,.34,'E4EEF1')
text(s,'Prefer validated radial geometry; use supported wall directions when arc coverage is insufficient.',.88,6.60,11.51,.23,13.5,DARK,True)
note(s,'Handling Partial Rim Observations',
'部分观测下，碗的实例分割可能完整，但可用于几何拟合的顶部深度点只覆盖一段圆弧。当前方法优先使用通过检查的圆心径向；只有圆弧覆盖或占用角度区间不足时，才采用有多尺度壁面支持的局部方向。备用分支不利用已被拒绝的短弧圆心来决定方向。这里左图对应实时规划预览，圆弧约 103 度，生成了 8 个候选并完成七段规划；旁边的几何图来自同次调试归档中的独立新采样，不能当成同一帧。它们均不代表本次进行了实机抓取。',
['docs/圆碗径向抓取修改_20260928.md','debug_output/circle_wall_fallback_20260928/live_plan/result.json','debug_output/circle_wall_fallback_20260928/fresh_geometry/radial_grasps.svg'])

# 6: geometric principle
s=base('Placement depends on the tool-to-center offset','PRINCIPLE / TAUGHT PLACEMENT','Geometric explanation under fixed bowl geometry, consistent engagement, and no slip.')
img(s,'contact_geometry.png',.66,1.89,5.15,4.72,crop=(0,0,.502,0))
caption(s,'Illustrative rendering; not a measured placement trial',.78,6.62,5.1)
label(s,'BOWL CENTER AT RELEASE',6.14,2.02,6)
text(s,'c = p + R r',6.14,2.46,6.42,.72,33,DARK,True,font='Liberation Serif')
text(s,'Bowl center = tool position + rotated object offset',6.16,3.25,6.1,.44,15,GREY)
points(s,[('Teaching fixes the release pose','Store full position and orientation in consistent frames.'),('Radial grasping constrains the offset','The aim is lower variation in the bowl-to-tool relationship.'),('Residual error remains','Calibration, contact depth and slip still affect placement.')],6.14,4.02,6.40,gap=.92)
note(s,'Why Grasp Consistency Helps Placement',
'这个式子是汇报中解释设计原理的重点。释放时碗中心等于工具位置，加上经工具姿态旋转后的工具到碗中心偏移。示教固定工具位置和姿态，径向抓取以及固定手指身份尝试让相对偏移保持一致。对于相同圆碗，在固定夹持高度、深度且不滑动的理想情况下，不同方位抓取可以获得相近的工具坐标系中心偏移。实际中仍有深度误差、接触变形和标定偏差。目前程序没有实现基于每次估计偏移的在线释放补偿，这里解释的是设计依据。',
['docs/圆碗径向抓取修改_20260928.md','src/grasp_executor/grasp_executor/grasp.py','src/grasp_executor/config/place_zones.json'])

# 7: complete task feasibility
s=base('Select a grasp that supports the complete task','CONTRIBUTION 03 / TASK FEASIBILITY','Selected target overlay from a real task | 29 Sep 2026, 14:26:44 | The overlay is a target pose.')
img(s,'grasp_pose.png',.65,1.94,8.14,5.05)
label(s,'SEVEN MOTION SEGMENTS',9.07,1.97,3.6)
segments=['Approach','Descend to grasp','Lift','Transfer to pre-release','Descend to release','Retreat','Return to saved home']
for i,v in enumerate(segments):
 text(s,f'{i+1:02d}',9.09,2.49+i*.43,.40,.28,12,BLUE,True)
 text(s,v,9.60,2.43+i*.43,3.07,.38,16,DARK)
text(s,'Try another candidate before motion if complete planning fails.',9.10,5.80,3.52,.77,17,DARK,True)
note(s,'Select Grasps for the Entire Task',
'局部几何合格只是第一步。程序对同一只碗生成有限数量的候选，并检查完整七段抓放返程路径。候选在尚未运动时不可规划，就可以尝试下一个；一旦动作开始，失败不会自动从张爪步骤重新抓取。返回起点使用保存的关节配置，避免相同末端位姿对应不同关节分支。右图叠加的是应用抓取偏移后的目标位姿，execution_started 只表示已经开始执行相关候选，并不能作为实测到位或抓牢的证明。',
['src/grasp_executor/grasp_executor/candidates.py','src/grasp_executor/grasp_executor/moveit_ros.py','debug_output/place_execution/20260929_142644/round_0001/grasp_pose_overlay.png'])

# 8: measured state and fallback
s=base('Advance execution using measured state','SYSTEM DESIGN / FEEDBACK','Gripper feedback: real task log, 28 Sep 2026 | TF fallback: isolated simulation, 29 Sep 2026')
img(s,'left_finger.jpg',.65,1.85,3.85,4.86,crop=(.40,.40,.20,.39))
label(s,'GRIPPER STATE CONFIRMATION',4.83,1.95,7.7)
text(s,'Confirm open before descent.\nConfirm closed and settled before lift.\nConfirm release before retreat.',4.83,2.38,7.60,1.32,21,DARK,True)
rect(s,4.82,3.93,7.79,.69,'E4EEF1')
text(s,'Recorded feedback: Close 679 → 37  |  Open 24 → 680',5.01,4.14,7.40,.30,16,BLUE,True)
label(s,'TEMPORARY TF LOSS',4.84,4.95,7.5)
text(s,'Use fresh joint measurements with validated forward kinematics.',4.84,5.34,7.56,.76,20,DARK,True)
text(s,'16 integration checks passed; 12 joint configurations cross-checked.\nGripper state confirmation does not verify that an object is held.',4.84,6.26,7.55,.53,13.5,GREY)
note(s,'Execution Advances on Measured State',
'夹爪部分从命令和固定等待，改为校验控制寄存器、开口位置及稳定时间。示例日志中，闭合从 679 变为 37，释放打开从 24 变为 680；这证明反馈状态发生变化，不直接证明碗被夹牢。机械臂正常使用 TF，在 TF 暂时不可用时使用新鲜六关节测量和已交叉核验的本地正运动学。启动核验、反馈新鲜度和到位条件继续生效。TF 故障注入结果来自隔离模拟控制器，没有连接机器人硬件，不能把模拟恢复当作已经完成所有实机故障验证。',
['src/grasp_executor/grasp_executor/gripper_command.py','debug_output/place_execution/20260928_211216/round_0003/grasp_executor.log','docs/TF反馈恢复_20260929.md','debug_output/tf_feedback_recovery_20260929/validation_summary.json'])

# 9: actual replay counts
s=base('More observations produced valid grasp candidates','EXPERIMENT / OFFLINE REPLAY','Historical replay evidence | circle_partial_arc_fix_20260928 and circle_wall_fallback_20260928')
img(s,'radial_geometry.png',.65,1.90,3.50,3.50)
img(s,'fallback_geometry.png',4.22,1.90,3.50,3.50)
caption(s,'Observed-contact radial candidates',.87,5.41,3.2)
caption(s,'Supported local-wall fallback',4.47,5.41,3.2)
label(s,'FRAMES WITH VALID CANDIDATES',8.12,2.02,4.6)
text(s,'14 / 20  →  18 / 20',8.12,2.59,4.53,.57,28,BLUE,True)
text(s,'Partial-arc acceptance',8.14,3.22,4.5,.42,17,GREY)
rect(s,8.14,3.85,4.47,.013,LIGHT)
text(s,'23 / 27  →  27 / 27',8.12,4.21,4.53,.57,28,TEAL,True)
text(s,'Supported local-wall fallback',8.14,4.84,4.5,.60,17,GREY)
rect(s,.70,6.14,11.91,.59,'E4EEF1')
text(s,'Two different replay sets. Candidate generation is not physical grasp success.',.92,6.31,11.43,.30,16,DARK,True)
note(s,'Offline Replay Results',
'第一组有 20 份历史点云，圆弧覆盖条件调整后，生成有效候选的帧数从 14 增加到 18。第二组有 27 份点云，在加入有支持的局部碗壁备用分支后，从 23 增加到 27。这是两个不同数据集、不同阶段的回放结果，不能把两个数列直接串成同一组实验。两侧是该项目实际回放可视化，目的是说明候选来源；数字以保存的 JSON 汇总为依据。这些结果支持候选生成可用性的改善，但不是抓取成功率，也没有证明当前版本在新场景中百分之百可用。',
['debug_output/circle_partial_arc_fix_20260928/replay.json','debug_output/circle_wall_fallback_20260928/replay.json'])

# 10: real trials and current features
s=base('A complete single-bowl demonstration workflow','IMPLEMENTED CAPABILITIES / HARDWARE RECORDS','Images from separate trials, 29 Sep 2026 | Completed command sequences; physical outcomes not automatically verified.')
for name,x,head,tm in [('blue_pose.png',.65,'BLUE TARGET','14:20:46'),('upper_left_pose.png',4.90,'UPPER LEFT','14:22:52'),('grasp_pose.png',9.15,'LOWER RIGHT','14:26:44')]:
 img(s,name,x,2.02,3.53,2.43)
 label(s,head,x,4.58,3.5)
 text(s,tm+' | command sequence completed',x,4.94,3.55,.57,12.5,GREY)
rect(s,.65,5.58,12.02,.013,LIGHT)
cols=[(.65,'One command, one bowl','Fresh observation; release, retreat and fixed home return.'),(4.90,'Six taught destinations','Five spatial zones plus blue; explicit XYZ input is also supported.'),(9.15,'Traceable results','Saved masks, candidate poses, point clouds and configuration snapshots.')]
for x,h,b in cols:
 text(s,h,x,5.87,3.65,.42,18,DARK,True);text(s,b,x,6.36,3.59,.48,14,GREY)
note(s,'Hardware Demonstration and Implemented Capabilities',
'这里展示三个独立任务保存的原始抓取位姿图，分别是 blue、左上和右下，图像并非同一次运动的连续帧。对应运行记录均有完整命令流程完成状态，但 grasp_verified 和 placement_verified 仍然为 false，所以不能直接把它们计作已独立验证的物理成功。功能方面，每次启动重新观察一只碗，按选择的示教目标完成放置、退回和返程。视觉增加桌面 ROI、分块推理及有限重观察，任务保存配置、候选和诊断图，便于逐次复现。',
['debug_output/place_execution/20260929_142046/result.json','debug_output/place_execution/20260929_142252/result.json','debug_output/place_execution/20260929_142644/result.json','New_run/README.md'])

# 11: future study on original table image
s=base('Measure placement repeatability directly','NEXT EXPERIMENT / PLANNED EVALUATION','Original scene image; no measured placement-error distribution is available in the current evidence set.')
img(s,'raw_table.png',.65,1.95,7.62,4.77)
label(s,'PLANNED COMPARISON',8.64,2.03,4)
text(s,'Wall-normal grasping\nvs. signed radial grasping',8.64,2.44,4.02,.89,21,DARK,True)
points(s,[('Vary the source position','Repeat across positions and grasp directions.'),('Measure physical outcomes','Grasp success, XY placement error and cycle time.'),('Separate the methods','Report radial and fallback trials independently.')],8.64,3.66,4.03,gap=.94)
rect(s,.89,6.11,6.93,.48,DARK)
text(s,'Target quantity: bowl-center distance to the taught target',1.07,6.24,6.55,.24,13,WHITE)
note(s,'Quantifying Placement Repeatability',
'下一步实验应直接验证设计假设：在相同碗、相同目标位置和相同夹持高度深度下，径向方法是否比原壁面方向方法具有更小的放置误差。改变碗的起始位置和抓取方位，进行重复试验，分别记录检测失败、无候选、规划失败、夹空、掉落和放偏。落点指标使用碗中心到目标中心的平面距离，报告均值、离散程度和最大值。图是实际桌面观测，但当前没有系统测量的落点分布，因此不展示虚构误差图。',
['docs/圆碗径向抓取修改_20260928.md','debug_output/table_roi_update_20260928/after/rgb.png'])

# 12: strong, evidence-bounded conclusion
s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(DARK)
img(s,'lifted_white.jpg',7.15,0,6.19,7.5,'cover')
rect(s,0,0,7.32,7.5,DARK)
label(s,'DESIGN CONTRIBUTIONS',.65,.55,6,'8ACFE1')
text(s,'Make the grasp relationship\nmore consistent.',.65,1.18,6.40,1.23,31,WHITE,True)
for yy,n,h,b in [(2.88,'01','Consistent geometry','Radial alignment and fixed finger identity.'),(3.92,'02','Observed contact support','Global shape with local wall validation.'),(4.96,'03','Complete-task feasibility','Select candidates for the full motion cycle.')]:
 text(s,n,.67,yy,.47,.41,16,'8ACFE1',True)
 text(s,h,1.27,yy-.03,5.62,.41,21,WHITE,True)
 text(s,b,1.27,yy+.43,5.62,.40,15,'C8DCE3')
rect(s,.68,6.20,5.93,.018,'446673')
text(s,'NEXT MILESTONE',.68,6.53,5.90,.23,10.5,'8ACFE1',True)
text(s,'Quantified physical success and placement repeatability',.68,6.91,6.13,.34,16,WHITE,True)
rect(s,8.45,6.71,4.36,.45,DARK)
text(s,'Original hardware photograph | 23 Sep 2026',8.57,6.84,4.13,.22,9.5,WHITE)
note(s,'Design Contributions and Next Milestone',
'总结本项目的三项设计贡献：第一，用有符号径向和固定内外手指约束抓取关系；第二，用整体形状指导方向，同时用真实局部壁面验证接触；第三，以完整抓放返程可行性选择候选。这些是相对于项目原实现的设计改进，并不在本汇报中声称已完成学术首创性证明。反馈和诊断改进使方法能够在现有硬件上实施与复查。下一阶段要通过对比实验量化物理抓取成功和放置重复性。',
['docs/圆碗径向抓取修改_20260928.md','docs/demo_implementation_20260925.md','docs/TF反馈恢复_20260929.md'])

pptx=OUT/'Bowl_Manipulation_Research_Report_EN_CN.pptx'
prs.save(pptx)
(OUT/'slide_notes.json').write_text(json.dumps(records,ensure_ascii=False,indent=2))
(OUT/'Chinese_Speaker_Notes.txt').write_text('\n\n'.join(f"第 {r['slide']} 页 | {r['title']}\n\n{r['notes_zh']}\n\n资料来源：\n"+'\n'.join(r['sources']) for r in records),encoding='utf-8')
# Verify editable slide text, notes, and physical page bounds.
errors=[]
for i,slide in enumerate(prs.slides,1):
 for sh in slide.shapes:
  if sh.has_text_frame and re.search(r'[\u4e00-\u9fff]',sh.text):errors.append(f'CJK on slide {i}')
  if sh.left < -200 or sh.top < -200 or sh.left+sh.width > prs.slide_width+15000 or sh.top+sh.height > prs.slide_height+15000:
   errors.append(f'out of bounds slide {i}: {getattr(sh,"text", "image")[:50]}')
 if not re.search(r'[\u4e00-\u9fff]',slide.notes_slide.notes_text_frame.text):errors.append(f'missing notes {i}')
report={'slides':len(prs.slides),'chinese_note_pages':len(records),'english_slide_text':True,'errors':errors,'image_generation':'built-in image_gen; 1 illustration reused on slides 2 and 6','hardware_commands_sent':False,'evidence_period':'2026-09-23 to 2026-09-29'}
(OUT/'validation.json').write_text(json.dumps(report,indent=2))
print(pptx)
print(json.dumps(report,ensure_ascii=False))
if errors:raise SystemExit(1)
